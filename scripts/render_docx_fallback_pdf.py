"""Create a visual-QA PDF from a DOCX when the canonical renderer is unavailable.

This is a diagnostic fallback, not a substitute deliverable.  It reads the
actual DOCX block order and styles with python-docx, then lays out the content
on Letter pages using bundled ReportLab.
"""

from __future__ import annotations

import argparse
import html
import re
from pathlib import Path
from typing import Iterable, Iterator, List, Sequence, Tuple, Union

from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    LongTable,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    TableStyle,
)


NAVY = colors.HexColor("#17365D")
BLUE = colors.HexColor("#1F4E78")
PALE_GRAY = colors.HexColor("#F2F2F2")
MID_GRAY = colors.HexColor("#D9E1F2")


def _register_fonts() -> Tuple[str, str, str]:
    candidates = (
        (
            Path(r"C:\Windows\Fonts\arial.ttf"),
            Path(r"C:\Windows\Fonts\arialbd.ttf"),
            Path(r"C:\Windows\Fonts\cour.ttf"),
        ),
        (
            Path(r"C:\Windows\Fonts\calibri.ttf"),
            Path(r"C:\Windows\Fonts\calibrib.ttf"),
            Path(r"C:\Windows\Fonts\consola.ttf"),
        ),
    )
    for regular, bold, mono in candidates:
        if regular.is_file() and bold.is_file() and mono.is_file():
            pdfmetrics.registerFont(TTFont("QaBody", str(regular)))
            pdfmetrics.registerFont(TTFont("QaBodyBold", str(bold)))
            pdfmetrics.registerFont(TTFont("QaMono", str(mono)))
            return "QaBody", "QaBodyBold", "QaMono"
    return "Helvetica", "Helvetica-Bold", "Courier"


BODY_FONT, BOLD_FONT, MONO_FONT = _register_fonts()


def _iter_blocks(parent: DocumentType) -> Iterator[Union[DocxParagraph, DocxTable]]:
    body = parent.element.body
    paragraph_type = parent.paragraphs[0]._p.__class__ if parent.paragraphs else None
    table_type = parent.tables[0]._tbl.__class__ if parent.tables else None
    for child in body.iterchildren():
        if paragraph_type is not None and isinstance(child, paragraph_type):
            yield DocxParagraph(child, parent)
        elif table_type is not None and isinstance(child, table_type):
            yield DocxTable(child, parent)


def _has_page_break(paragraph: DocxParagraph) -> bool:
    return bool(paragraph._p.xpath('.//w:br[@w:type="page"]'))


def _safe_markup(text: str) -> str:
    escaped = html.escape(text.replace("`", ""), quote=False)

    def split_token(match: re.Match[str]) -> str:
        token = match.group(0)
        return "&#8203;".join(token[index : index + 18] for index in range(0, len(token), 18))

    return re.sub(r"[^\s]{28,}", split_token, escaped)


def _styles():
    sample = getSampleStyleSheet()
    return {
        "body": ParagraphStyle(
            "QaBody",
            parent=sample["BodyText"],
            fontName=BODY_FONT,
            fontSize=9.2,
            leading=11.3,
            textColor=colors.HexColor("#202020"),
            spaceAfter=5,
            allowWidows=0,
            allowOrphans=0,
        ),
        "title": ParagraphStyle(
            "QaTitle",
            parent=sample["Title"],
            fontName=BOLD_FONT,
            fontSize=22,
            leading=26,
            alignment=TA_CENTER,
            textColor=NAVY,
            spaceAfter=13,
            keepWithNext=1,
        ),
        "h1": ParagraphStyle(
            "QaH1",
            parent=sample["Heading1"],
            fontName=BOLD_FONT,
            fontSize=15,
            leading=18,
            textColor=NAVY,
            spaceBefore=11,
            spaceAfter=5,
            keepWithNext=1,
        ),
        "h2": ParagraphStyle(
            "QaH2",
            parent=sample["Heading2"],
            fontName=BOLD_FONT,
            fontSize=12,
            leading=14,
            textColor=BLUE,
            spaceBefore=9,
            spaceAfter=4,
            keepWithNext=1,
        ),
        "h3": ParagraphStyle(
            "QaH3",
            parent=sample["Heading3"],
            fontName=BOLD_FONT,
            fontSize=10.2,
            leading=12.2,
            textColor=BLUE,
            spaceBefore=7,
            spaceAfter=3,
            keepWithNext=1,
        ),
        "bullet": ParagraphStyle(
            "QaBullet",
            parent=sample["BodyText"],
            fontName=BODY_FONT,
            fontSize=9.2,
            leading=11.3,
            leftIndent=14,
            firstLineIndent=-8,
            bulletIndent=4,
            spaceAfter=2.5,
        ),
        "cell": ParagraphStyle(
            "QaCell",
            parent=sample["BodyText"],
            fontName=BODY_FONT,
            fontSize=7.1,
            leading=8.6,
            textColor=colors.HexColor("#202020"),
        ),
        "cell_header": ParagraphStyle(
            "QaCellHeader",
            parent=sample["BodyText"],
            fontName=BOLD_FONT,
            fontSize=7.1,
            leading=8.6,
            textColor=colors.white,
        ),
    }


def _column_widths(rows: Sequence[Sequence[str]], available: float) -> List[float]:
    count = len(rows[0])
    scores = []
    for index in range(count):
        longest = max(len(str(row[index])) for row in rows)
        average = sum(min(len(str(row[index])), 80) for row in rows) / len(rows)
        scores.append(max(9.0, min(35.0, 0.65 * longest + 0.35 * average)))
    total = sum(scores)
    return [available * score / total for score in scores]


def _table_flowable(table: DocxTable, styles, available_width: float) -> LongTable:
    raw_rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    data: List[List[Paragraph]] = []
    for row_index, row in enumerate(raw_rows):
        style = styles["cell_header"] if row_index == 0 else styles["cell"]
        data.append([Paragraph(_safe_markup(value), style) for value in row])
    widths = _column_widths(raw_rows, available_width)
    result = LongTable(data, colWidths=widths, repeatRows=1, splitByRow=1, hAlign="CENTER")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), BLUE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#A6A6A6")),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for row_index in range(2, len(data), 2):
        commands.append(("BACKGROUND", (0, row_index), (-1, row_index), PALE_GRAY))
    result.setStyle(TableStyle(commands))
    return result


def _header_footer(canvas, document) -> None:
    canvas.saveState()
    width, height = letter
    canvas.setFont(BODY_FONT, 7.5)
    canvas.setFillColor(colors.HexColor("#595959"))
    canvas.drawRightString(width - 0.72 * inch, height - 0.37 * inch, "REP-02  /  Oturum 5")
    canvas.drawString(0.72 * inch, 0.33 * inch, "D3 sentetik test ve kanıt kapanışı")
    # Keep the page label comfortably inside the raster crop used for visual QA.
    canvas.drawRightString(width - 0.90 * inch, 0.33 * inch, "Sayfa " + str(document.page))
    canvas.restoreState()


def render(docx_path: Path, pdf_path: Path) -> None:
    source = Document(docx_path)
    styles = _styles()
    available_width = letter[0] - 1.44 * inch
    story: List[object] = []
    list_number = 0
    for block in _iter_blocks(source):
        if isinstance(block, DocxTable):
            story.append(_table_flowable(block, styles, available_width))
            story.append(Spacer(1, 5))
            continue
        if _has_page_break(block):
            story.append(PageBreak())
            continue
        text = block.text.strip()
        if not text:
            continue
        style_name = block.style.name if block.style is not None else "Normal"
        if style_name == "Title":
            style = styles["title"]
            list_number = 0
        elif style_name == "Heading 1":
            style = styles["h1"]
            list_number = 0
        elif style_name == "Heading 2":
            style = styles["h2"]
            list_number = 0
        elif style_name == "Heading 3":
            style = styles["h3"]
            list_number = 0
        elif style_name.startswith("List Bullet"):
            style = styles["bullet"]
            text = "• " + text
            list_number = 0
        elif style_name.startswith("List Number"):
            list_number += 1
            style = styles["bullet"]
            text = str(list_number) + ". " + text
        else:
            style = styles["body"]
            list_number = 0
        story.append(Paragraph(_safe_markup(text), style))

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    document = BaseDocTemplate(
        str(pdf_path),
        pagesize=letter,
        leftMargin=0.72 * inch,
        rightMargin=0.72 * inch,
        topMargin=0.62 * inch,
        bottomMargin=0.57 * inch,
        title="REP-02 D3 closeout visual QA preview",
        author="OpenAI Codex",
    )
    frame = Frame(
        document.leftMargin,
        document.bottomMargin,
        document.width,
        document.height,
        id="normal",
        leftPadding=0,
        rightPadding=0,
        topPadding=8,
        bottomPadding=8,
    )
    document.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=_header_footer)])
    document.build(story)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    render(args.input, args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
