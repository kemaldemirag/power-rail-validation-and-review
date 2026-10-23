"""Render the REP-02 D3 closeout Markdown as a styled DOCX deliverable.

This helper intentionally supports only the Markdown constructs used by the
closeout report: headings, paragraphs, bullets, numbered items, pipe tables,
inline code, and explicit ``<!-- PAGEBREAK -->`` markers.
"""

from __future__ import annotations

import argparse
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, List, Optional, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "17365D"
BLUE = "1F4E78"
PALE_BLUE = "D9EAF7"
PALE_GRAY = "F2F2F2"
MID_GRAY = "D9E1F2"
WHITE = "FFFFFF"
BODY_FONT = "Aptos"
CODE_FONT = "Consolas"


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, *, top: int = 70, start: int = 85, bottom: int = 70, end: int = 85) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _keep_with_next(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    properties.append(keep)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Sayfa ")
    run.font.name = BODY_FONT
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


INLINE_CODE = re.compile(r"(`[^`]+`)")


def _add_inline(
    paragraph,
    text: str,
    *,
    bold: bool = False,
    color: Optional[RGBColor] = None,
    size: float = 10,
) -> None:
    for part in INLINE_CODE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(max(8.5, size - 1))
            run.font.color.rgb = RGBColor(31, 78, 121)
        else:
            run = paragraph.add_run(part)
            run.font.name = BODY_FONT
            run.font.size = Pt(size)
            run.bold = bold
            if color is not None:
                run.font.color.rgb = color


def _split_table_row(line: str) -> List[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_table_separator(line: str) -> bool:
    cells = _split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _set_table_geometry(table, column_count: int) -> None:
    """Pin table width/indent/grid so layout is deterministic across renderers."""

    # Text width is 10,166 DXA. Reserve the 85-DXA cell inset as table indent
    # so the visible border and the surrounding paragraph text share an edge.
    total_width = 10081
    widths = [total_width // column_count] * column_count
    widths[-1] += total_width - sum(widths)

    properties = table._tbl.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.insert(0, table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(total_width))

    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "85")

    layout = properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_column.set(qn("w:w"), str(width))

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(width))


def _add_table(document: Document, header: Sequence[str], rows: Sequence[Sequence[str]]) -> None:
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    heading_row = table.rows[0]
    _repeat_table_header(heading_row)
    for index, value in enumerate(header):
        cell = heading_row.cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_shading(cell, BLUE)
        _set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        _add_inline(paragraph, value, bold=True, color=RGBColor(255, 255, 255), size=8.5)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            _set_cell_margins(cell)
            if row_index % 2:
                _set_cell_shading(cell, PALE_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            _add_inline(paragraph, value, size=8.5)
            for run in paragraph.runs:
                if run.font.name != CODE_FONT:
                    run.font.size = Pt(8.5)
    _set_table_geometry(table, len(header))
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _configure_document(document: Document) -> None:
    for section in document.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        section.header_distance = Inches(0.28)
        section.footer_distance = Inches(0.28)

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header.add_run("REP-02  /  Oturum 5")
        header_run.font.name = BODY_FONT
        header_run.font.size = Pt(8)
        header_run.font.color.rgb = RGBColor(89, 89, 89)

        footer = section.footer.paragraphs[0]
        footer_run = footer.add_run("D3 sentetik test ve kanıt kapanışı  •  ")
        footer_run.font.name = BODY_FONT
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = RGBColor(89, 89, 89)
        _add_page_number(footer)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color, before, after in (
        ("Title", 25, NAVY, 0, 12),
        ("Heading 1", 18, NAVY, 14, 6),
        ("Heading 2", 14, BLUE, 12, 5),
        ("Heading 3", 11.5, BLUE, 9, 4),
    ):
        style = document.styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ("List Bullet", "List Number"):
        style = document.styles[list_style]
        style.font.name = BODY_FONT
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(2.5)

    properties = document.core_properties
    properties.title = "REP-02 Oturum 5 — Regülatör İncelemesi, Dosya API ve D3 Kapanışı"
    properties.subject = "WF-02-05, WF-02-06 ve D3 sentetik test/kanıt teslimi"
    properties.author = "OpenAI Codex — REP-02 çalışma kaydı"
    properties.keywords = "REP-02, regulator review, evidence, file API, D3"
    properties.comments = "Sentetik yazılım kapsamı; canlı entegrasyon veya donanım onayı değildir."
    properties.created = datetime(2026, 10, 19, 12, 0, tzinfo=timezone.utc)
    properties.modified = datetime(2026, 10, 19, 12, 0, tzinfo=timezone.utc)


def _flush_paragraph(document: Document, lines: List[str]) -> None:
    if not lines:
        return
    paragraph = document.add_paragraph()
    _add_inline(paragraph, " ".join(line.strip() for line in lines))
    lines.clear()


def render_markdown(lines: Iterable[str], output_path: Path) -> None:
    document = Document()
    _configure_document(document)
    material = list(lines)
    paragraph_lines: List[str] = []
    index = 0
    first_heading = True
    while index < len(material):
        line = material[index].rstrip("\n")
        stripped = line.strip()
        if not stripped:
            _flush_paragraph(document, paragraph_lines)
            index += 1
            continue
        if stripped == "<!-- PAGEBREAK -->":
            _flush_paragraph(document, paragraph_lines)
            document.add_page_break()
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(material) and _is_table_separator(material[index + 1].strip()):
            _flush_paragraph(document, paragraph_lines)
            header = _split_table_row(stripped)
            rows: List[List[str]] = []
            index += 2
            while index < len(material) and material[index].strip().startswith("|"):
                values = _split_table_row(material[index].strip())
                if len(values) == len(header):
                    rows.append(values)
                index += 1
            _add_table(document, header, rows)
            continue
        heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading_match:
            _flush_paragraph(document, paragraph_lines)
            level = len(heading_match.group(1))
            text_value = heading_match.group(2)
            if first_heading and level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_inline(paragraph, text_value, bold=True, color=RGBColor.from_string(NAVY), size=25)
                first_heading = False
            else:
                paragraph = document.add_heading(level=level)
                heading_size = {1: 18, 2: 14, 3: 11.5}[level]
                _add_inline(
                    paragraph,
                    text_value,
                    bold=True,
                    color=RGBColor.from_string(NAVY if level == 1 else BLUE),
                    size=heading_size,
                )
                _keep_with_next(paragraph)
            index += 1
            continue
        if stripped.startswith("- "):
            _flush_paragraph(document, paragraph_lines)
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline(paragraph, stripped[2:])
            index += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            _flush_paragraph(document, paragraph_lines)
            paragraph = document.add_paragraph(style="List Number")
            _add_inline(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            index += 1
            continue
        paragraph_lines.append(stripped)
        index += 1
    _flush_paragraph(document, paragraph_lines)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    render_markdown(args.input.read_text(encoding="utf-8").splitlines(), args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
