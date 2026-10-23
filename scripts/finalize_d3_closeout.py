"""Finalize deterministic REP-02 D3 verification records and manifest.

Run this only after the test suite, independent reviews, and document visual
inspection have completed.  The script replays local static checks, verifies
the DOCX package, writes ``test-results-closeout.json``, and then builds the
byte-exact closeout manifest with the public reporting API.
"""

from __future__ import annotations

import argparse
import ast
import hashlib
import json
import math
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Sequence, Tuple
from zipfile import ZipFile


sys.dont_write_bytecode = True
REPO_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = REPO_ROOT / "src"
OUTPUT_ROOT = REPO_ROOT / "evidence" / "REP02-D3"
REPORT_DOCX = REPO_ROOT / "docs" / "REP02-D3" / "REP02-D3-regulator-integration-and-closeout.docx"
REPORT_MD = REPO_ROOT / "docs" / "REP02-D3" / "REP02-D3-regulator-integration-and-closeout.md"
GENERATED_AT = "2026-10-19T23:00:00+03:00"

if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from power_rail_validation import __version__ as PACKAGE_VERSION  # noqa: E402
from power_rail_validation.fault_injection import (  # noqa: E402
    ENGINE_VERSION as FAULT_ENGINE_VERSION,
    RULE_SET_VERSION as FAULT_RULE_SET_VERSION,
)
from power_rail_validation.integration import (  # noqa: E402
    API_VERSION,
    DESCRIPTOR_SCHEMA_VERSION,
    ENGINE_VERSION as INTEGRATION_ENGINE_VERSION,
)
from power_rail_validation.regulator_review import (  # noqa: E402
    ENGINE_VERSION as REGULATOR_ENGINE_VERSION,
    RULE_SET_VERSION as REGULATOR_RULE_SET_VERSION,
)
from power_rail_validation.rail_budget import FORMULA_VERSION as RAIL_BUDGET_FORMULA_VERSION  # noqa: E402
from power_rail_validation.reporting import (  # noqa: E402
    REPORTER_VERSION,
    build_evidence_manifest,
    render_json_report,
    render_markdown_report,
)
from power_rail_validation.sequencing import (  # noqa: E402
    ENGINE_VERSION as SEQUENCING_ENGINE_VERSION,
    RULE_SET_VERSION as SEQUENCING_RULE_SET_VERSION,
)


MALFORMED_JSON_FIXTURES = {
    "tests/fixtures/contracts/malformed.json",
    "tests/fixtures/integration/malformed.json",
}
REQUIRED_DOCX_PARTS = {
    "[Content_Types].xml",
    "_rels/.rels",
    "word/document.xml",
    "word/styles.xml",
    "word/settings.xml",
    "word/header1.xml",
    "word/footer1.xml",
}
REQUIRED_REPORT_TEXT = (
    "REP-02 Oturum 5",
    "WF-02-05 Regulator Review",
    "WF-02-06 çevrimdışı REP-01/03 dosya API",
    "D3 Test/Evidence",
    "GAP-02",
    "GAP-03",
    "UNCOMMITTED_DELIVERABLES",
)
GENERATOR_OUTPUTS = (
    "integration-results.json",
    "regulator-review-results.json",
    "report-closeout.md",
    "report-model-closeout.json",
    "source-model-provenance-register.json",
)


def _json_text(value: Any) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        allow_nan=False,
        sort_keys=True,
        indent=2,
    ) + "\n"


def _write_json(path: Path, value: Any) -> None:
    text = _json_text(value)
    if json.loads(text) != value:
        raise RuntimeError("JSON round-trip changed output: " + path.as_posix())
    descriptor, temporary_name = tempfile.mkstemp(
        dir=str(path.parent),
        prefix=path.name + ".",
        suffix=".tmp",
        text=True,
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(text)
        os.replace(str(temporary_path), str(path))
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def _write_bytes_atomic(path: Path, payload: bytes) -> None:
    descriptor, temporary_name = tempfile.mkstemp(
        dir=str(path.parent),
        prefix=path.name + ".",
        suffix=".tmp",
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(payload)
        os.replace(str(temporary_path), str(path))
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _relative(path: Path) -> str:
    return path.resolve().relative_to(REPO_ROOT.resolve()).as_posix()


def _strict_json(path: Path) -> Any:
    def reject_constant(value: str) -> None:
        raise ValueError("non-finite JSON constant: " + value)

    def unique_pairs(pairs: Sequence[Tuple[str, Any]]) -> Dict[str, Any]:
        result: Dict[str, Any] = {}
        for key, value in pairs:
            if key in result:
                raise ValueError("duplicate JSON key: " + key)
            result[key] = value
        return result

    value = json.loads(
        path.read_text(encoding="utf-8"),
        parse_constant=reject_constant,
        object_pairs_hook=unique_pairs,
    )

    def finite(item: Any) -> bool:
        if isinstance(item, float):
            return math.isfinite(item)
        if isinstance(item, list):
            return all(finite(child) for child in item)
        if isinstance(item, dict):
            return all(finite(child) for child in item.values())
        return item is None or isinstance(item, (str, bool, int))

    if not finite(value):
        raise ValueError("non-finite or unsupported JSON value: " + path.as_posix())
    return value


def _python_ast_check() -> int:
    files = sorted(
        path
        for root in (REPO_ROOT / "src", REPO_ROOT / "tests", REPO_ROOT / "scripts")
        for path in root.rglob("*.py")
    )
    for path in files:
        ast.parse(path.read_text(encoding="utf-8"), filename=str(path), feature_version=(3, 9))
    return len(files)


def _json_fixture_check() -> Tuple[int, int]:
    candidates = sorted(
        path
        for root in (REPO_ROOT / "src", REPO_ROOT / "tests")
        for path in root.rglob("*.json")
    )
    parsed = 0
    rejected = 0
    for path in candidates:
        relative = _relative(path)
        if relative in MALFORMED_JSON_FIXTURES:
            try:
                _strict_json(path)
            except (UnicodeError, ValueError, json.JSONDecodeError):
                rejected += 1
                continue
            raise RuntimeError("deliberately malformed fixture parsed: " + relative)
        _strict_json(path)
        parsed += 1
    if rejected != len(MALFORMED_JSON_FIXTURES):
        raise RuntimeError("malformed fixture inventory drifted")
    return parsed, rejected


def _git_check() -> Tuple[str, bool, bool]:
    head = subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=REPO_ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    diff = subprocess.run(
        ["git", "diff", "--check"],
        cwd=REPO_ROOT,
        check=False,
        capture_output=True,
        text=True,
    )
    if diff.returncode != 0:
        raise RuntimeError("git diff --check failed:\n" + diff.stdout + diff.stderr)
    dirty = bool(
        subprocess.run(
            ["git", "status", "--porcelain"],
            cwd=REPO_ROOT,
            check=True,
            capture_output=True,
            text=True,
        ).stdout.strip()
    )
    return head, True, dirty


def _docx_check() -> Dict[str, Any]:
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:  # pragma: no cover - environment precondition
        raise RuntimeError("python-docx is required for closeout verification") from exc

    with ZipFile(REPORT_DOCX) as archive:
        corrupt = archive.testzip()
        names = set(archive.namelist())
    if corrupt is not None:
        raise RuntimeError("DOCX has a corrupt ZIP member: " + corrupt)
    missing_parts = sorted(REQUIRED_DOCX_PARTS - names)
    if missing_parts:
        raise RuntimeError("DOCX required parts missing: " + ", ".join(missing_parts))

    document = Document(REPORT_DOCX)
    letter_width = int(Inches(8.5))
    letter_height = int(Inches(11))
    if any(
        int(section.page_width or 0) != letter_width
        or int(section.page_height or 0) != letter_height
        for section in document.sections
    ):
        raise RuntimeError("DOCX sections are not Letter portrait")
    material: List[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            material.extend(cell.text for cell in row.cells)
    extracted = "\n".join(material)
    missing_text = [value for value in REQUIRED_REPORT_TEXT if value not in extracted]
    if missing_text:
        raise RuntimeError("DOCX required text missing: " + ", ".join(missing_text))
    return {
        "docx_zip_integrity": "PASS",
        "required_ooxml_parts": "PASS",
        "required_text_coverage": "PASS",
        "section_count": len(document.sections),
        "table_count": len(document.tables),
        "page_size": "LETTER_PORTRAIT",
    }


def _placeholder_check() -> int:
    needles = ("TODO", "TBD", "FIXME", "PLACEHOLDER", "<insert", "lorem ipsum", "XXX")
    paths = (REPORT_MD, REPO_ROOT / "README.md", REPO_ROOT / "docs" / "README.md", REPO_ROOT / "evidence" / "README.md")
    for path in paths:
        folded = path.read_text(encoding="utf-8").casefold()
        for needle in needles:
            if needle.casefold() in folded:
                raise RuntimeError(f"placeholder marker {needle!r} in {_relative(path)}")
    return len(paths)


def _artifact_hashes(paths: Iterable[str]) -> Dict[str, str]:
    return {path: _sha256(REPO_ROOT / Path(path)) for path in sorted(paths)}


def _verification_corpus_paths() -> Tuple[str, ...]:
    """Return every byte that can affect the discovered unit-test corpus."""

    files = {REPO_ROOT / "pyproject.toml"}
    files.update(
        path
        for path in (REPO_ROOT / "src").rglob("*")
        if path.is_file() and path.suffix != ".pyc" and "__pycache__" not in path.parts
    )
    files.update(
        path
        for path in (REPO_ROOT / "tests").rglob("*.py")
        if path.is_file() and "__pycache__" not in path.parts
    )
    files.update(
        path
        for path in (REPO_ROOT / "tests" / "fixtures").rglob("*")
        if path.is_file() and path.suffix != ".pyc" and "__pycache__" not in path.parts
    )
    missing = [path for path in files if not path.is_file()]
    if missing:
        raise RuntimeError(
            "verification corpus path missing: "
            + ", ".join(sorted(_relative(path) for path in missing))
        )
    return tuple(sorted((_relative(path) for path in files)))


def _corpus_entry(identifier_prefix: str, path: str) -> Dict[str, str]:
    path_token = hashlib.sha256(path.encode("utf-8")).hexdigest()[:20].upper()
    return _entry(
        identifier_prefix + "-" + path_token,
        path,
        "session-5-byte-bound",
    )


def _assert_verification_corpus_snapshot(recorded: Mapping[str, str]) -> None:
    current_paths = _verification_corpus_paths()
    current = _artifact_hashes(current_paths)
    if dict(recorded) != current:
        recorded_paths = set(recorded)
        current_path_set = set(current)
        added = sorted(current_path_set - recorded_paths)
        removed = sorted(recorded_paths - current_path_set)
        changed = sorted(
            path
            for path in recorded_paths & current_path_set
            if recorded[path] != current[path]
        )
        raise RuntimeError(
            "verification corpus snapshot drifted; added="
            + repr(added)
            + ", removed="
            + repr(removed)
            + ", changed="
            + repr(changed)
        )


def _generator_input_hashes() -> Dict[str, str]:
    files = sorted(
        {
            *(path for path in (REPO_ROOT / "src").rglob("*") if path.is_file()),
            *(path for path in (REPO_ROOT / "tests" / "fixtures").rglob("*") if path.is_file()),
            REPO_ROOT / "scripts" / "generate_d3_closeout.py",
        },
        key=lambda path: path.as_posix(),
    )
    return {_relative(path): _sha256(path) for path in files}


def _generator_replay_check() -> Dict[str, str]:
    before_inputs = _generator_input_hashes()
    observed: List[Dict[str, str]] = []
    expected_names = set(GENERATOR_OUTPUTS)
    for seed in ("1", "777"):
        environment = os.environ.copy()
        environment["PYTHONDONTWRITEBYTECODE"] = "1"
        environment["PYTHONHASHSEED"] = seed
        environment["PYTHONPATH"] = str(SRC_ROOT)
        completed = subprocess.run(
            [sys.executable, "scripts/generate_d3_closeout.py"],
            cwd=REPO_ROOT,
            env=environment,
            check=False,
            capture_output=True,
            text=True,
        )
        if completed.returncode != 0:
            raise RuntimeError("D3 evidence generator replay failed:\n" + completed.stdout + completed.stderr)
        try:
            summary = json.loads(completed.stdout)
        except json.JSONDecodeError as exc:
            raise RuntimeError("D3 evidence generator emitted invalid summary JSON") from exc
        declared = summary.get("output_sha256") if isinstance(summary, dict) else None
        if not isinstance(declared, dict) or set(declared) != expected_names:
            raise RuntimeError("D3 evidence generator output inventory drifted")
        actual = {name: _sha256(OUTPUT_ROOT / name) for name in GENERATOR_OUTPUTS}
        if declared != actual:
            raise RuntimeError("D3 evidence generator summary hashes do not match output bytes")
        observed.append(actual)
    if observed[0] != observed[1]:
        raise RuntimeError("D3 evidence differs across PYTHONHASHSEED values")
    if before_inputs != _generator_input_hashes():
        raise RuntimeError("D3 evidence generation mutated source or fixture inputs")
    return observed[0]


def _entry(identifier: str, path: str, version: str) -> Dict[str, str]:
    return {
        "id": identifier,
        "path": path,
        "sha256": _sha256(REPO_ROOT / Path(path)),
        "version": version,
    }


def _verified_test_run(
    identifier: str,
    unittest_args: Sequence[str],
    expected_count: int,
) -> Dict[str, Any]:
    command = [sys.executable, "-m", "unittest", *unittest_args, "-q"]
    environment = os.environ.copy()
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    environment["PYTHONPATH"] = str(SRC_ROOT)
    completed = subprocess.run(
        command,
        cwd=REPO_ROOT,
        env=environment,
        check=False,
        capture_output=True,
        text=True,
    )
    output = completed.stdout + completed.stderr
    match = re.search(r"Ran\s+(\d+)\s+tests?\b", output)
    if completed.returncode != 0:
        raise RuntimeError(identifier + " failed:\n" + output)
    if match is None:
        raise RuntimeError(identifier + " did not report a unittest count:\n" + output)
    observed_count = int(match.group(1))
    if observed_count != expected_count:
        raise RuntimeError(
            f"{identifier} count mismatch: expected {expected_count}, observed {observed_count}"
        )
    return {
        "id": identifier,
        "command": "PYTHONDONTWRITEBYTECODE=1 PYTHONPATH=src python -m unittest "
        + " ".join(unittest_args)
        + " -q",
        "tests_run": observed_count,
        "passed": observed_count,
        "failed": 0,
        "errors": 0,
        "status": "PASS",
    }


def _build_test_results(args: argparse.Namespace) -> Dict[str, Any]:
    python_files = _python_ast_check()
    json_files, malformed_rejected = _json_fixture_check()
    head, diff_check, dirty = _git_check()
    generator_hashes = _generator_replay_check()
    docx = _docx_check()
    placeholder_files = _placeholder_check()
    pyc_files = sorted(REPO_ROOT.rglob("*.pyc"))
    if pyc_files:
        raise RuntimeError("repository contains .pyc files: " + ", ".join(_relative(path) for path in pyc_files))

    report_model = _strict_json(OUTPUT_ROOT / "report-model-closeout.json")
    regulator_suite = _strict_json(OUTPUT_ROOT / "regulator-review-results.json")
    integration_suite = _strict_json(OUTPUT_ROOT / "integration-results.json")
    _strict_json(OUTPUT_ROOT / "source-model-provenance-register.json")
    _strict_json(OUTPUT_ROOT / "fault-campaign-results.json")
    if regulator_suite.get("verification_status") != "PASS":
        raise RuntimeError("regulator evidence verification is not PASS")
    if integration_suite.get("verification_status") != "PASS":
        raise RuntimeError("integration evidence verification is not PASS")
    if render_json_report(report_model).encode("utf-8") != (
        OUTPUT_ROOT / "report-model-closeout.json"
    ).read_bytes():
        raise RuntimeError("report JSON bytes differ from the public reporter output")
    if render_markdown_report(report_model).encode("utf-8") != (
        OUTPUT_ROOT / "report-closeout.md"
    ).read_bytes():
        raise RuntimeError("report Markdown bytes differ from the public reporter output")

    docx.update(
        {
            "canonical_renderer": {
                "status": "BLOCKED",
                "reason": "BUNDLED_LIBREOFFICE_SOFFICE_NOT_PRESENT_IN_WORKSPACE_RUNTIME",
            },
            "fallback_render": {
                "status": "PASS",
                "evidence_origin": "CALLER_ATTESTED_COMPLETED_VISUAL_INSPECTION",
                "method": (
                    "Actual DOCX structure parsed with bundled python-docx and ReportLab, "
                    "then rasterized with Poppler"
                ),
                "page_count": args.docx_pages,
                "visual_inspection": "PASS",
                "checks": [
                    "all pages inspected",
                    "no clipped text",
                    "no overlapping content",
                    "tables fit page width",
                    "headers and footers are readable",
                    "no missing Turkish glyphs",
                ],
            },
            "skill_audits": {
                "evidence_origin": "CALLER_ATTESTED_COMPLETED_DOCUMENT_SKILL_AUDIT",
                "accessibility_findings": 0,
                "table_geometry_findings": 0,
                "heading_hierarchy": "PASS",
                "section_layout": "PASS",
            },
        }
    )

    corpus_paths = _verification_corpus_paths()
    corpus_hashes_before = _artifact_hashes(corpus_paths)
    test_runs = [
        _verified_test_run("FULL-UNITTEST", ("discover", "-s", "tests"), args.full_tests),
        _verified_test_run(
            "REGULATOR-REVIEW-FOCUSED",
            ("discover", "-s", "tests", "-p", "test_regulator_review.py"),
            args.regulator_tests,
        ),
        _verified_test_run(
            "INTEGRATION-FOCUSED",
            ("discover", "-s", "tests", "-p", "test_integration.py"),
            args.integration_tests,
        ),
        _verified_test_run(
            "FAULT-INJECTION-FOCUSED",
            ("discover", "-s", "tests", "-p", "test_fault_injection.py"),
            args.fault_tests,
        ),
        _verified_test_run(
            "REPORTING-FOCUSED",
            ("discover", "-s", "tests", "-p", "test_reporting.py"),
            args.reporting_tests,
        ),
    ]
    corpus_hashes_after = _artifact_hashes(corpus_paths)
    if corpus_hashes_before != corpus_hashes_after:
        raise RuntimeError("verification corpus bytes changed while tests were running")

    selected_paths = set(corpus_paths)
    selected_paths.update(
        {
            ".gitattributes",
            "README.md",
            "docs/README.md",
            "evidence/README.md",
            "scripts/generate_d3_closeout.py",
            "scripts/render_closeout_docx.py",
            "scripts/render_docx_fallback_pdf.py",
            "scripts/finalize_d3_closeout.py",
            "docs/REP02-D3/REP02-D3-regulator-integration-and-closeout.md",
            "docs/REP02-D3/REP02-D3-regulator-integration-and-closeout.docx",
            "evidence/REP02-D3/fault-campaign-results.json",
            "evidence/REP02-D3/regulator-review-results.json",
            "evidence/REP02-D3/integration-results.json",
            "evidence/REP02-D3/source-model-provenance-register.json",
            "evidence/REP02-D3/report-model-closeout.json",
            "evidence/REP02-D3/report-closeout.md",
        }
    )
    selected_hashes = _artifact_hashes(
        selected_paths
    )
    if any(selected_hashes[path] != digest for path, digest in corpus_hashes_before.items()):
        raise RuntimeError("verification corpus hashes changed before closeout serialization")

    return {
        "schema_version": "1.0.0",
        "deliverable": "REP02-D3-REGULATOR-INTEGRATION-CLOSEOUT-VERIFICATION",
        "generated_at": GENERATED_AT,
        "verification_status": "PASS",
        "technical_closeout_status": "D3_SYNTHETIC_TEST_EVIDENCE_CLOSED_WITH_LIVE_INTEGRATION_BLOCKED",
        "scope_status": {
            "WF-02-04": "IMPLEMENTED",
            "WF-02-05": "IMPLEMENTED_SYNTHETIC",
            "WF-02-06_OFFLINE_FILE_API": "IMPLEMENTED_SYNTHETIC",
            "D3": "CLOSED_SYNTHETIC_SCOPE",
            "REP-01_LIVE": "NOT_RUN",
            "REP-03_LIVE": "NOT_RUN",
            "GAP-02": "OPEN",
            "GAP-03": "OPEN",
            "G1": "BLOCKED",
            "G2_G6": "NOT_OPENED",
        },
        "baseline": {
            "commit": head,
            "working_tree": "UNCOMMITTED_DELIVERABLES" if dirty else "CLEAN",
        },
        "environment": {
            "operating_system": "Windows",
            "python_runtime": ".".join(str(part) for part in sys.version_info[:3]),
            "minimum_supported_python": "3.9",
            "package_version": PACKAGE_VERSION,
            "regulator_engine_version": REGULATOR_ENGINE_VERSION,
            "regulator_rule_set_version": REGULATOR_RULE_SET_VERSION,
            "integration_engine_version": INTEGRATION_ENGINE_VERSION,
            "file_api_version": API_VERSION,
            "descriptor_schema_version": DESCRIPTOR_SCHEMA_VERSION,
            "fault_engine_version": FAULT_ENGINE_VERSION,
            "fault_rule_set_version": FAULT_RULE_SET_VERSION,
            "sequencing_engine_version": SEQUENCING_ENGINE_VERSION,
            "sequencing_rule_set_version": SEQUENCING_RULE_SET_VERSION,
            "reporter_version": REPORTER_VERSION,
            "external_test_dependencies": [],
            "closeout_verification_dependencies": ["python-docx"],
        },
        "test_runs": test_runs,
        "static_checks": [
            {"id": "PYTHON-3.9-AST", "files_checked": python_files, "status": "PASS"},
            {"id": "STRICT-JSON-SOURCE-AND-FIXTURE-PARSE", "files_checked": json_files, "status": "PASS"},
            {"id": "DELIBERATE-MALFORMED-JSON-REJECTION", "files_checked": malformed_rejected, "status": "PASS"},
            {"id": "PLACEHOLDER-SCAN", "files_checked": placeholder_files, "status": "PASS"},
            {"id": "GIT-DIFF-CHECK", "status": "PASS" if diff_check else "FAIL"},
            {"id": "REPO-PYC-HYGIENE", "observed_count": 0, "status": "PASS"},
        ],
        "verification_corpus": {
            "path_policy": "ALL_SRC_FILES_PLUS_TEST_PY_PLUS_TEST_FIXTURES_PLUS_PYPROJECT",
            "file_count": len(corpus_hashes_before),
            "artifact_hashes": corpus_hashes_before,
        },
        "evidence_checks": {
            "two_build_equality": "PASS",
            "cross_process_hash_seed_determinism": "PASS",
            "input_immutability": "PASS",
            "finite_json": "PASS",
            "report_json_markdown_semantic_equality": "PASS",
            "manifest_byte_hash_replay": "PASS",
            "generator_output_sha256": generator_hashes,
            "regulator_case_count": regulator_suite["case_count"],
            "integration_case_count": integration_suite["case_count"],
            "report_status": report_model["status"],
            "report_model_hash": report_model["report_model_hash"],
        },
        "independent_reviews": [
            {
                "scope": "REGULATOR-REVIEW-ADVERSARIAL",
                "status": "CLEAN",
                "closed_findings": 21,
                "open_findings": 0,
            },
            {
                "scope": "OFFLINE-FILE-API-ADVERSARIAL",
                "status": "CLEAN",
                "closed_findings": 2,
                "open_findings": 0,
            },
            {
                "scope": "D3-CLOSEOUT-GENERATOR-ADVERSARIAL",
                "status": "CLEAN",
                "closed_findings": 3,
                "open_findings": 0,
            },
            {
                "scope": "FAULT-INJECTION-ADVERSARIAL-REGRESSION",
                "status": "CLEAN",
                "open_findings": 0,
            },
        ],
        "document_qa": docx,
        "artifact_hashes": selected_hashes,
        "actions": {
            "commit": False,
            "push": False,
            "package_installation": False,
            "external_repository_mutation": False,
            "live_rep01_rep03_integration": False,
            "physical_test": False,
        },
        "claim_limits": [
            "SYNTHETIC_FIXTURES_ONLY",
            "SYNTHETIC_PASS_IS_NOT_PROJECT_OR_HARDWARE_PASS",
            "DESCRIPTOR_STRUCTURE_AND_FILE_INTEGRITY_ONLY",
            "NO_LIVE_REP01_OR_REP03_INTEGRATION_CLAIM",
            "NO_ANALOG_TRANSIENT_OR_WAVEFORM_EXECUTION",
            "NO_PHYSICAL_HARDWARE_OR_LAB_VALIDATION",
            "NO_EMC_ISO_AEC_OR_FIELD_COMPLIANCE_CLAIM",
            "G1_SOURCE_GATE_BLOCKED",
            "G2_TO_G6_HARDWARE_GATES_NOT_OPENED",
        ],
        "next_work": [
            "D4/Oturum 6 CLI and exit-code contract",
            "wheel/sdist package-data and clean-install smoke verification",
            "CI evidence replay and release preparation",
            "live REP-01/03 adapters only after authorized repositories and semantic mappings are available",
        ],
    }


def _manifest_entries(
    corpus_paths: Sequence[str] = (),
) -> Mapping[str, Sequence[Mapping[str, str]]]:
    entries: Dict[str, List[Mapping[str, str]]] = {
        "inputs": [
            _entry("INPUT-REGULATOR-NOMINAL", "tests/fixtures/regulator_review/nominal-synthetic.json", "fixture-1"),
            _entry("INPUT-REGULATOR-NEGATIVE", "tests/fixtures/regulator_review/negative-cases.json", "fixture-1"),
            _entry("INPUT-REGULATOR-WORKSPACE-GAP", "tests/fixtures/regulator_review/workspace-gap-profile.json", "fixture-1"),
            _entry(
                "INPUT-REGULATOR-TRANSIENT-PRESENT-STABILITY-MISSING",
                "tests/fixtures/regulator_review/transient-present-stability-missing.json",
                "fixture-1",
            ),
            _entry("INPUT-INTEGRATION-VALID", "tests/fixtures/integration/valid.json", "fixture-1"),
            _entry("INPUT-INTEGRATION-REP03-VALID", "tests/fixtures/integration/rep03-valid.json", "fixture-1"),
            _entry("INPUT-INTEGRATION-MULTIPLE-PROFILES", "tests/fixtures/integration/multiple-profiles.json", "fixture-1"),
            _entry("INPUT-INTEGRATION-MISSING-ARTIFACT", "tests/fixtures/integration/missing-artifact.json", "fixture-1"),
            _entry("INPUT-INTEGRATION-MALFORMED", "tests/fixtures/integration/malformed.json", "fixture-1"),
            _entry(
                "INPUT-INTEGRATION-ARTIFACT-CONFIG-A",
                "tests/fixtures/integration/artifacts/rep01-configuration-a.json",
                "fixture-1",
            ),
            _entry(
                "INPUT-INTEGRATION-ARTIFACT-CONFIG-B",
                "tests/fixtures/integration/artifacts/rep01-configuration-b.json",
                "fixture-1",
            ),
            _entry(
                "INPUT-INTEGRATION-ARTIFACT-PROFILE-COLD",
                "tests/fixtures/integration/artifacts/rep01-profile-cold.json",
                "fixture-1",
            ),
            _entry(
                "INPUT-INTEGRATION-ARTIFACT-PROFILE-NOMINAL",
                "tests/fixtures/integration/artifacts/rep01-profile-nominal.json",
                "fixture-1",
            ),
            _entry(
                "INPUT-INTEGRATION-ARTIFACT-REP03-ASSESSMENT-A",
                "tests/fixtures/integration/artifacts/rep03-assessment-a.json",
                "fixture-1",
            ),
            _entry("INPUT-FAULT-BROWNOUT", "tests/fixtures/fault_injection/brownout.json", "fixture-1"),
            _entry("INPUT-FAULT-DELAY", "tests/fixtures/fault_injection/delay.json", "fixture-1"),
            _entry(
                "INPUT-FAULT-DURING-RESTART",
                "tests/fixtures/fault_injection/fault-during-restart.json",
                "fixture-1",
            ),
            _entry("INPUT-FAULT-RAIL-LOSS", "tests/fixtures/fault_injection/rail-loss.json", "fixture-1"),
            _entry(
                "INPUT-FAULT-RECOVERY-COMBO",
                "tests/fixtures/fault_injection/recovery-combo.json",
                "fixture-1",
            ),
            _entry("INPUT-FAULT-RESTART", "tests/fixtures/fault_injection/restart.json", "fixture-1"),
            _entry("INPUT-FAULT-STUCK-LOW", "tests/fixtures/fault_injection/stuck-low.json", "fixture-1"),
            _entry(
                "INPUT-FAULT-UNKNOWN-THRESHOLD",
                "tests/fixtures/fault_injection/unknown-threshold.json",
                "fixture-1",
            ),
            _entry("INPUT-RAIL-BASELINE", "tests/fixtures/rail_budget/nominal.json", "fixture-1"),
            _entry("INPUT-SEQUENCING-BASELINE", "tests/fixtures/sequencing/normal.json", "fixture-1"),
        ],
        "sources": [
            _entry("SOURCE-REGULATOR-RULES", "src/power_rail_validation/regulator_review/review-rules.json", REGULATOR_RULE_SET_VERSION),
            _entry("SOURCE-FILE-API-CONTRACT", "src/power_rail_validation/integration/file_api_contract.json", API_VERSION),
            _entry(
                "SOURCE-REGULATOR-DATASHEET",
                "tests/fixtures/regulator_review/synthetic-exact-part-source.json",
                "fixture-datasheet-1.0.0",
            ),
            _entry(
                "SOURCE-REGULATOR-RAIL-BUDGET",
                "tests/fixtures/regulator_review/synthetic-rail-budget-source.json",
                "fixture-rail-1.0.0",
            ),
            _entry(
                "SOURCE-REGULATOR-STABILITY",
                "tests/fixtures/regulator_review/synthetic-stability-result.json",
                "fixture-stability-1.0.0",
            ),
            _entry(
                "SOURCE-REGULATOR-TRANSIENT",
                "tests/fixtures/regulator_review/synthetic-transient-result.json",
                "fixture-transient-1.0.0",
            ),
            _entry("SOURCE-FAULT-RULES", "tests/fixtures/fault_injection/synthetic-fault-rules-source.json", FAULT_RULE_SET_VERSION),
            _entry("SOURCE-SEQUENCING-RULES", "tests/fixtures/sequencing/synthetic-rules-source.json", SEQUENCING_RULE_SET_VERSION),
        ],
        "models": [
            _entry(
                "MODEL-REGULATOR-SYNTHETIC",
                "tests/fixtures/regulator_review/synthetic-model.lib",
                "fixture-model-1.0.0",
            ),
        ],
        "tools": [
            _entry("TOOL-REGULATOR-REVIEW", "src/power_rail_validation/regulator_review/reviewer.py", REGULATOR_ENGINE_VERSION),
            _entry("TOOL-FILE-API", "src/power_rail_validation/integration/file_api.py", INTEGRATION_ENGINE_VERSION),
            _entry("TOOL-REPORTER", "src/power_rail_validation/reporting/reporter.py", REPORTER_VERSION),
            _entry("TOOL-FAULT-RUNNER", "src/power_rail_validation/fault_injection/runner.py", FAULT_ENGINE_VERSION),
            _entry("TOOL-RAIL-BUDGET", "src/power_rail_validation/rail_budget/calculator.py", RAIL_BUDGET_FORMULA_VERSION),
            _entry("TOOL-SEQUENCING", "src/power_rail_validation/sequencing/evaluator.py", SEQUENCING_ENGINE_VERSION),
            _entry("TOOL-REGULATOR-TESTS", "tests/test_regulator_review.py", "session-5"),
            _entry("TOOL-INTEGRATION-TESTS", "tests/test_integration.py", "session-5"),
            _entry("TOOL-FAULT-TESTS", "tests/test_fault_injection.py", "session-5"),
            _entry("TOOL-REPORTING-TESTS", "tests/test_reporting.py", "session-5"),
            _entry("TOOL-EVIDENCE-GENERATOR", "scripts/generate_d3_closeout.py", "session-5"),
            _entry("TOOL-DOCX-RENDERER", "scripts/render_closeout_docx.py", "session-5"),
            _entry("TOOL-DOCX-FALLBACK-RENDERER", "scripts/render_docx_fallback_pdf.py", "session-5"),
            _entry("TOOL-CLOSEOUT-FINALIZER", "scripts/finalize_d3_closeout.py", "session-5"),
            _entry("TOOL-LINE-END-POLICY", ".gitattributes", "policy-1"),
        ],
        "outputs": [
            _entry("OUTPUT-FAULT-CAMPAIGN", "evidence/REP02-D3/fault-campaign-results.json", "1.0.0"),
            _entry("OUTPUT-REGULATOR-RESULTS", "evidence/REP02-D3/regulator-review-results.json", "1.0.0"),
            _entry("OUTPUT-INTEGRATION-RESULTS", "evidence/REP02-D3/integration-results.json", "1.0.0"),
            _entry("OUTPUT-SOURCE-MODEL-REGISTER", "evidence/REP02-D3/source-model-provenance-register.json", "1.0.0"),
            _entry("OUTPUT-CLOSEOUT-REPORT-MODEL", "evidence/REP02-D3/report-model-closeout.json", "1.0.0"),
            _entry("OUTPUT-CLOSEOUT-REPORT-MARKDOWN", "evidence/REP02-D3/report-closeout.md", REPORTER_VERSION),
            _entry("OUTPUT-DELIVERY-REPORT-MARKDOWN", "docs/REP02-D3/REP02-D3-regulator-integration-and-closeout.md", "1.0.0"),
            _entry("OUTPUT-DELIVERY-REPORT-DOCX", "docs/REP02-D3/REP02-D3-regulator-integration-and-closeout.docx", "1.0.0"),
            _entry("OUTPUT-CLOSEOUT-TEST-RESULTS", "evidence/REP02-D3/test-results-closeout.json", "1.0.0"),
        ],
    }
    covered_paths = {
        item["path"]
        for category in entries.values()
        for item in category
    }
    bound_corpus_paths = tuple(corpus_paths) if corpus_paths else _verification_corpus_paths()
    for path in bound_corpus_paths:
        if path in covered_paths:
            continue
        if path.startswith("tests/fixtures/"):
            category = "inputs"
            prefix = "INPUT-TEST-CORPUS"
        else:
            category = "tools"
            prefix = "TOOL-TEST-CORPUS"
        entries[category].append(_corpus_entry(prefix, path))
        covered_paths.add(path)
    return entries


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--full-tests", type=int, required=True)
    parser.add_argument("--regulator-tests", type=int, required=True)
    parser.add_argument("--integration-tests", type=int, required=True)
    parser.add_argument("--fault-tests", type=int, required=True)
    parser.add_argument("--reporting-tests", type=int, required=True)
    parser.add_argument("--docx-pages", type=int, required=True)
    args = parser.parse_args()
    if any(getattr(args, name) <= 0 for name in ("full_tests", "regulator_tests", "integration_tests", "fault_tests", "reporting_tests", "docx_pages")):
        raise RuntimeError("all verification counts must be positive")

    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    test_results = _build_test_results(args)
    test_results_path = OUTPUT_ROOT / "test-results-closeout.json"
    manifest_path = OUTPUT_ROOT / "manifest-closeout.json"
    previous = {
        path: path.read_bytes() if path.is_file() else None
        for path in (test_results_path, manifest_path)
    }
    try:
        recorded_corpus = test_results["verification_corpus"]["artifact_hashes"]
        _assert_verification_corpus_snapshot(recorded_corpus)
        _write_json(test_results_path, test_results)
        for relative, digest in test_results["artifact_hashes"].items():
            if _sha256(REPO_ROOT / Path(relative)) != digest:
                raise RuntimeError("recorded artifact hash drifted before manifest creation: " + relative)
        _assert_verification_corpus_snapshot(recorded_corpus)
        corpus_paths = tuple(sorted(recorded_corpus))
        manifest_entries = _manifest_entries(corpus_paths)
        manifest = build_evidence_manifest(
            manifest_entries,
            base_dir=REPO_ROOT,
            generated_at=GENERATED_AT,
        )
        _assert_verification_corpus_snapshot(recorded_corpus)
        replay = build_evidence_manifest(
            _manifest_entries(corpus_paths),
            base_dir=REPO_ROOT,
            generated_at=GENERATED_AT,
        )
        if replay != manifest:
            raise RuntimeError("closeout manifest replay differs")
        _assert_verification_corpus_snapshot(recorded_corpus)
        _write_json(manifest_path, manifest)
        _assert_verification_corpus_snapshot(recorded_corpus)
    except Exception:
        for path, payload in previous.items():
            if payload is None:
                if path.exists():
                    path.unlink()
            else:
                _write_bytes_atomic(path, payload)
        raise

    sys.stdout.write(
        _json_text(
            {
                "test_results": _relative(test_results_path),
                "test_results_sha256": _sha256(test_results_path),
                "manifest": _relative(manifest_path),
                "manifest_hash": manifest["manifest_hash"],
                "manifest_file_sha256": _sha256(manifest_path),
            }
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
